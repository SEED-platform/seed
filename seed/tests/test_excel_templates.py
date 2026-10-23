"""
SEED Platform (TM), Copyright (c) Alliance for Energy Innovation, LLC, and other contributors.
See also https://github.com/SEED-platform/seed/blob/main/LICENSE.md
"""

from contextlib import closing
from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.etree.ElementTree import tostring

import pandas as pd
from django.test import SimpleTestCase
from openpyxl import load_workbook

from seed.utils.excel import write_template_dataframe

TEMPLATE_DIR = Path(__file__).parents[1] / "utils"
REFERENCE_DIR = Path(__file__).parent / "data" / "styleframe_exports"
FACILITY_TEMPLATE = TEMPLATE_DIR / "CTS Facility Upload Template.xlsx"
PLAN_TEMPLATE = TEMPLATE_DIR / "facilities_plan_output.xlsx"


def facility_dataframe(populated):
    df = pd.read_excel(FACILITY_TEMPLATE, sheet_name="Facility Upload Template", dtype=object)
    columns = df.columns
    df.columns = df.loc[0]
    if populated:
        # Include rows beyond the template, missing values, leading zeros and Excel dates.
        for index in range(3):
            df.loc[3 + index] = pd.Series(
                {
                    "Facility Status": 1,
                    "Agency Designated Covered Facility ID": f"000{index + 1}",
                    "Government Owned Facility": index == 0,
                    "Date of Facility Re-/Retro-Commissioning": date(2026, 1, index + 1),
                    "Facility Name": f"Facility {index + 1}",
                    "Zip Code": "08001",
                    "Fiscal Year": 2026,
                    "Gross Square Footage": 3.5,
                    "Number of Buildings Metered for Electricity": 0,
                    "Covered Facility Internal Comments": '=HYPERLINK("https://example.com", "Details")',
                },
                dtype=object,
            )
    df.columns = columns
    return df


def plan_dataframes(populated):
    properties = pd.DataFrame()
    components = pd.DataFrame()
    if populated:
        properties = pd.DataFrame(
            {
                "Site Address": ["001 Test Street", "002 Test Street"],
                "Electricity": [10, 0],
                "Include": [True, False],
                "Total Energy Usage": [10.5, 0],
            }
        )
        components = pd.DataFrame(
            {
                "Site Address": ["001 Test Street", "002 Test Street"],
                "Uniformat Category": ["D3010", "D3020"],
                "RSL": [5.5, 0],
                "Replacement Cost": [1000, 0],
                "Component_SubType": ["Heat pump", None],
                "Capacity_cooling": [0, None],
            }
        )
    return {
        "How To Guide": pd.read_excel(PLAN_TEMPLATE, sheet_name="How To Guide"),
        "1. Facilities Plan": properties,
        "2. Prioritized Bldg Components": components,
        "3. ECM Data Dictionary": pd.read_excel(PLAN_TEMPLATE, sheet_name="3. ECM Data Dictionary"),
    }


class ExcelTemplateTests(SimpleTestCase):
    def assert_matches_styleframe(self, output, reference):
        # Compare resolved styles, not StyleFrame's randomized style names or ZIP timestamps.
        with closing(load_workbook(output)) as actual, closing(load_workbook(REFERENCE_DIR / reference)) as expected:
            assert actual.sheetnames == expected.sheetnames
            assert actual.loaded_theme == expected.loaded_theme
            assert actual.epoch == expected.epoch
            assert actual.defined_names == expected.defined_names
            for sheet, reference_sheet in zip(actual, expected, strict=True):
                with self.subTest(sheet=sheet.title):
                    assert sheet.calculate_dimension() == reference_sheet.calculate_dimension()
                    assert sheet.freeze_panes == reference_sheet.freeze_panes
                    assert sheet.sheet_state == reference_sheet.sheet_state
                    assert sheet.merged_cells == reference_sheet.merged_cells
                    assert sheet.print_area == reference_sheet.print_area
                    assert sheet.print_title_rows == reference_sheet.print_title_rows
                    assert sheet.print_title_cols == reference_sheet.print_title_cols
                    for attribute in (
                        "sheet_properties",
                        "sheet_format",
                        "views",
                        "auto_filter",
                        "protection",
                        "print_options",
                        "page_margins",
                        "page_setup",
                        "HeaderFooter",
                        "data_validations",
                    ):
                        assert tostring(getattr(sheet, attribute).to_tree()) == tostring(getattr(reference_sheet, attribute).to_tree())
                    for attribute in ("row_dimensions", "column_dimensions"):
                        dimensions = getattr(sheet, attribute)
                        reference_dimensions = getattr(reference_sheet, attribute)
                        assert dimensions.keys() == reference_dimensions.keys()
                        for key, dimension in dimensions.items():
                            assert dict(dimension) == dict(reference_dimensions[key])
                    for row, reference_row in zip(sheet, reference_sheet, strict=True):
                        for cell, reference_cell in zip(row, reference_row, strict=True):
                            with self.subTest(cell=cell.coordinate):
                                assert cell.value == reference_cell.value
                                assert cell.data_type == reference_cell.data_type
                                assert cell.number_format == reference_cell.number_format
                                assert cell.comment == reference_cell.comment
                                assert cell.hyperlink == reference_cell.hyperlink
                                for attribute in ("font", "fill", "border", "alignment", "protection"):
                                    assert copy(getattr(cell, attribute)) == copy(getattr(reference_cell, attribute))

    def test_facility_workbooks_match_styleframe(self):
        for populated in (False, True):
            with self.subTest(populated=populated):
                output = BytesIO()
                with pd.ExcelWriter(output, engine="openpyxl") as writer:
                    write_template_dataframe(
                        facility_dataframe(populated), writer, FACILITY_TEMPLATE, "Facility Upload Template", row_to_add_filters=1
                    )
                self.assert_matches_styleframe(output, f"facility_{'populated' if populated else 'empty'}.xlsx")

    def test_facilities_plan_workbooks_match_styleframe(self):
        for populated in (False, True):
            with self.subTest(populated=populated):
                output = BytesIO()
                with pd.ExcelWriter(output, engine="openpyxl") as writer:
                    for sheet_name, df in plan_dataframes(populated).items():
                        if sheet_name in ("How To Guide", "3. ECM Data Dictionary"):
                            write_template_dataframe(df, writer, PLAN_TEMPLATE, sheet_name, sheet_name=sheet_name)
                        else:
                            df.to_excel(writer, sheet_name=sheet_name, index=False)
                self.assert_matches_styleframe(output, f"plan_{'populated' if populated else 'empty'}.xlsx")
