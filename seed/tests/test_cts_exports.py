"""
SEED Platform (TM), Copyright (c) Alliance for Energy Innovation, LLC, and other contributors.
See also https://github.com/SEED-platform/seed/blob/main/LICENSE.md
"""

import tempfile
from contextlib import closing
from datetime import date
from io import BytesIO
from pathlib import Path

from buildingsync_asset_extractor.cts import building_sync_to_cts
from django.test import SimpleTestCase
from django.urls import reverse
from openpyxl import load_workbook

from seed.tests.util import AccessLevelBaseTestCase


class CTSFacilityExportTests(AccessLevelBaseTestCase):
    def test_export_preserves_mixed_values_and_groups_facilities(self):
        # String identifiers must retain leading zeros alongside numeric and boolean cells.
        cycle = self.cycle_factory.get_cycle(start=date(2026, 1, 1), end=date(2027, 1, 1))
        views = [
            self.property_view_factory.get_property_view(
                cycle=cycle,
                no_default_data=True,
                gross_floor_area=area,
                extra_data={
                    "Covered Building (75% Energy Usage)": "Yes",
                    "Sub-agency Acronym": "DOE",
                    "Installation Code": code,
                    "Government Owned": True,
                    "Energy Management System at Building's Installation": "50001Ready",
                    "Installation City": "Denver",
                    "Installation State": "CO",
                    "Installation Postal Code": "08001",
                },
            )
            for code, area in [("000123", 1000), ("000123", 2500), ("000456", 3000)]
        ]

        response = self.client.post(
            f"{reverse('api:v3:properties-facility-bps-export-to-cts')}?organization_id={self.org.id}",
            {"property_view_ids": [view.id for view in views]},
            content_type="application/json",
        )

        assert response.status_code == 200
        with closing(load_workbook(BytesIO(response.content))) as workbook:
            sheet = workbook.active
            headers = [cell.value for cell in sheet[2]]
            records = [dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=4, values_only=True)]
            facilities = {
                row["Agency Designated Covered Facility ID"]: row
                for row in records
                if row["Agency Designated Covered Facility ID"] is not None
            }

            assert set(facilities) == {"000123", "000456"}
            assert facilities["000123"]["Gross Square Footage"] == 3.5
            assert facilities["000456"]["Gross Square Footage"] == 3
            assert facilities["000123"]["Government Owned Facility"] is True
            assert facilities["000123"]["Energy Management System Status"] == 2
            assert facilities["000123"]["Fiscal Year"] == 2027
            assert facilities["000123"]["Zip Code"] == "08001"
            assert facilities["000123"]["Number of Buildings Metered for Electricity"] == 0
            assert facilities["000123"]["Date of Facility Re-/Retro-Commissioning"] is None

    def test_export_with_no_properties_preserves_template(self):
        response = self.client.post(
            f"{reverse('api:v3:properties-facility-bps-export-to-cts')}?organization_id={self.org.id}",
            {"property_view_ids": []},
            content_type="application/json",
        )

        assert response.status_code == 200
        with closing(load_workbook(BytesIO(response.content))) as workbook:
            assert workbook.active["A2"].value == "Facility Status"
            assert workbook.active["C4"].value is None


class CTSBuildingSyncExportTests(SimpleTestCase):
    def test_export_preserves_identification_and_numeric_values(self):
        # Exercise the upstream extractor's pandas 3 and openpyxl export path.
        xml = """<BuildingSync xmlns="http://buildingsync.net/schemas/bedes-auc/2019">
          <Facilities><Facility>
            <Sites><Site><Buildings><Building><FloorAreas><FloorArea>
              <FloorAreaType>Gross</FloorAreaType><FloorAreaValue>2500</FloorAreaValue>
            </FloorArea></FloorAreas></Building></Buildings></Site></Sites>
            <Reports><Report><RetrocommissioningAudit>true</RetrocommissioningAudit></Report></Reports>
            <UserDefinedFields>
              <UserDefinedField><FieldName>Agency Designated Covered Facility ID</FieldName><FieldValue>000123</FieldValue></UserDefinedField>
              <UserDefinedField><FieldName>Sub-agency Acronym</FieldName><FieldValue>DOE</FieldValue></UserDefinedField>
              <UserDefinedField><FieldName>Facility Name</FieldName><FieldValue>Test Facility</FieldValue></UserDefinedField>
            </UserDefinedFields>
          </Facility></Facilities>
        </BuildingSync>"""

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "facility.xml"
            output = Path(directory) / "evaluation.xlsx"
            source.write_text(xml, encoding="utf-8")
            building_sync_to_cts([source], output)

            with closing(load_workbook(output)) as workbook:
                sheet = workbook.active
                headers = [cell.value for cell in sheet[1]]
                records = [dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=2, values_only=True)]
                records = [row for row in records if row["Agency Designated Covered Facility ID"] == "000123"]

                assert len(records) == 1
                record = records[0]
                assert record["Agency Designated Covered Facility ID"] == "000123"
                assert record["Sub-agency Acronym"] == "DOE"
                assert record["Facility Name"] == "Test Facility"
                assert record["Gross Evaluated Square Footage"] == 2500
                assert record["Retro/Re-Commissioning Assessment"] is True
                assert record["Estimated Annual Energy Savings"] == 0
