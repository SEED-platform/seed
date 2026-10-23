"""
SEED Platform (TM), Copyright (c) Alliance for Energy Innovation, LLC, and other contributors.
See also https://github.com/SEED-platform/seed/blob/main/LICENSE.md
"""

from contextlib import closing
from itertools import chain
from pathlib import Path

import pandas as pd
from buildingsync_asset_extractor.cts.excel_template import CellStyle
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def write_template_dataframe(
    df: pd.DataFrame,
    writer: pd.ExcelWriter,
    template_path: Path,
    template_sheet_name: str,
    *,
    sheet_name: str = "Sheet1",
    row_to_add_filters: int | None = None,
) -> None:
    """Write a populated template DataFrame with the former StyleFrame export formatting.

    Like the extractor's template writer, retain only the formatting StyleFrame
    retained, not every setting in the source workbook. DataFrames must contain
    the template's original rows and columns, optionally with additional data.
    """
    template_rows, template_cols = pd.read_excel(template_path, sheet_name=template_sheet_name).shape
    df.to_excel(writer, sheet_name=sheet_name, index=False, na_rep="")
    sheet = writer.sheets[sheet_name]
    sheet.sheet_view.rightToLeft = False

    with closing(load_workbook(template_path)) as workbook:
        template = workbook[template_sheet_name]
        default_style = CellStyle()
        for col, (column_name, values) in enumerate(df.items(), start=1):
            for row, value in enumerate(chain([column_name], values), start=1):
                in_template = col <= template_cols and row <= template_rows + 1
                style = CellStyle.from_cell(template.cell(row=row, column=col)) if in_template else default_style
                if row > 1 and "=HYPERLINK" in str(value):
                    style = style._replace(font_color="000000FF", underline="single")
                style.apply_to(sheet.cell(row=row, column=col), value)

        for col in range(1, template_cols + 1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = template.column_dimensions[letter].width
        for row in range(1, template_rows + 2):
            sheet.row_dimensions[row].height = template.row_dimensions[row].height

    if row_to_add_filters is not None:
        filter_row = row_to_add_filters + 1
        sheet.auto_filter.ref = f"A{filter_row}:{get_column_letter(len(df.columns))}{filter_row}"
